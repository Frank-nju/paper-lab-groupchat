"""本地只读证据包：把赛题、论文和附件变成可定位的审计输入。"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterable, Iterator
from xml.etree import ElementTree


SCHEMA_VERSION = 1
_TEXT_SUFFIXES = {".txt", ".md", ".rst", ".json"}
_SUPPORTED_ATTACHMENT_SUFFIXES = {
    ".csv",
    ".tsv",
    ".json",
    ".xlsx",
    ".xlsm",
    ".pdf",
    ".txt",
    ".md",
    ".npy",
    ".npz",
    ".png",
    ".jpg",
    ".jpeg",
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _clean_text(value: str) -> str:
    value = value.replace("\u00a0", " ").replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"[ \t]+", " ", value).strip()


def _normal_text(value: str) -> str:
    return re.sub(r"\s+", "", value)


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    parts: list[str] = []
    for node in root.iter():
        if node.tag.endswith("}t") and node.text:
            parts.append(node.text)
        elif node.tag.endswith("}p"):
            parts.append("\n")
    return "".join(parts)


def _pdf_pages(path: Path) -> list[str]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency failure is environment-specific
        raise RuntimeError("解析 PDF 需要 pypdf") from exc
    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # pragma: no cover - malformed PDFs vary by file
            pages.append(f"[PDF 页面提取失败：{exc}]")
    return pages


def _document_pages(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf_pages(path)
    if suffix == ".docx":
        return [_docx_text(path)]
    return [_read_text(path)]


def _anchors(text: str) -> list[str]:
    values = re.findall(r"(?:图|表|公式|式)\s*[0-9一二三四五六七八九十百]+(?:[-.]\d+)?", text)
    return sorted(set(value.replace(" ", "") for value in values))


def _page_records(path: Path, source_id: str) -> list[dict[str, Any]]:
    pages = _document_pages(path)
    records: list[dict[str, Any]] = []
    for page_number, raw in enumerate(pages, start=1):
        text = _clean_text(raw)
        normalized_length = len(_normal_text(text))
        status = "ok" if normalized_length >= 40 else "sparse"
        records.append(
            {
                "source_id": source_id,
                "page": page_number,
                "location": f"{path.name} p.{page_number}",
                "text": text,
                "char_count": len(text),
                "normalized_char_count": normalized_length,
                "status": status,
                "anchors": _anchors(text),
            }
        )
    return records


def _chunks(pages: Iterable[dict[str, Any]], source_id: str) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    for page in pages:
        text = page["text"]
        if not text:
            continue
        size = 1600
        overlap = 180
        start = 0
        while start < len(text):
            end = min(len(text), start + size)
            if end < len(text):
                boundary = max(text.rfind("\n", start, end), text.rfind("。", start, end))
                if boundary > start + 400:
                    end = boundary + 1
            fragment = text[start:end].strip()
            if fragment:
                stable = hashlib.sha1(
                    f"{source_id}:{page['page']}:{start}:{fragment}".encode("utf-8")
                ).hexdigest()[:12]
                chunks.append(
                    {
                        "evidence_id": f"E-{stable}",
                        "source_id": source_id,
                        "location": page["location"],
                        "page_start": page["page"],
                        "page_end": page["page"],
                        "char_start": start,
                        "char_end": end,
                        "anchors": _anchors(fragment),
                        "text": fragment,
                    }
                )
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)
    return chunks


_ZH_NUMBERS = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}


def _question_number(value: str, fallback: int) -> int:
    value = value.strip()
    if value.isdigit():
        return int(value)
    if value in _ZH_NUMBERS:
        return _ZH_NUMBERS[value]
    if value.startswith("十"):
        return 10 + _ZH_NUMBERS.get(value[1:], 0)
    if "十" in value:
        left, right = value.split("十", 1)
        return _ZH_NUMBERS.get(left, 1) * 10 + _ZH_NUMBERS.get(right, 0)
    return fallback


def _question_index(problem_path: Path, pages: list[dict[str, Any]]) -> list[dict[str, Any]]:
    text = "\n".join(page["text"] for page in pages)
    pattern = re.compile(r"问题\s*([一二三四五六七八九十0-9]+)|第\s*([一二三四五六七八九十0-9]+)\s*问")
    matches = list(pattern.finditer(text))
    questions: list[dict[str, Any]] = []
    for index, match in enumerate(matches, start=1):
        number = _question_number(match.group(1) or match.group(2), index)
        start = match.start()
        end = matches[index].start() if index < len(matches) else len(text)
        content = text[start:end].strip()
        questions.append(
            {
                "question_id": f"Q{number}",
                "title": match.group(0).strip(),
                "text": content,
                "location": f"{problem_path.name} chars {start}-{end}",
                "source_id": pages[0]["source_id"] if pages else "",
            }
        )
    return questions


def _tokens(text: str) -> set[str]:
    values = re.findall(r"[\u4e00-\u9fff]{2,}|[A-Za-z][A-Za-z0-9_]{1,}|\d+(?:\.\d+)?", text)
    return {value.lower() for value in values if len(value) > 1}


def _retrieve(question: dict[str, Any], chunks: list[dict[str, Any]], limit: int = 6) -> list[dict[str, Any]]:
    query = _tokens(question["text"])
    ranked: list[tuple[int, dict[str, Any]]] = []
    for chunk in chunks:
        overlap = len(query & _tokens(chunk["text"]))
        ranked.append((overlap, chunk))
    ranked.sort(key=lambda item: (-item[0], item[1]["evidence_id"]))
    return [chunk for score, chunk in ranked[:limit] if score > 0] or chunks[: min(limit, len(chunks))]


def _csv_summary(path: Path, source_id: str, source_uri: str) -> dict[str, Any]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
        reader = csv.reader(stream, delimiter=delimiter)
        rows = list(reader)
    headers = rows[0] if rows else []
    body = rows[1:] if rows else []
    missing = {header or f"column_{i+1}": 0 for i, header in enumerate(headers)}
    for row in body:
        for i, header in enumerate(headers):
            if i >= len(row) or not str(row[i]).strip():
                missing[header or f"column_{i+1}"] += 1
    return {
        "source_id": source_id,
        "source_uri": source_uri,
        "format": path.suffix.lower().lstrip("."),
        "status": "ok",
        "rows": len(body),
        "columns": headers,
        "missing_by_column": missing,
        "preview": body[:5],
    }


def _xlsx_summary(path: Path, source_id: str, source_uri: str) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {
            "source_id": source_id,
            "source_uri": source_uri,
            "format": path.suffix.lower().lstrip("."),
            "status": "blocked",
            "error": "缺少 openpyxl，不能可靠审计 Excel 附件",
        }
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        iterator = worksheet.iter_rows(values_only=True)
        first = next(iterator, ())
        headers = [str(value) if value is not None else f"column_{i+1}" for i, value in enumerate(first)]
        rows = 0
        missing = {header: 0 for header in headers}
        numeric: dict[str, list[float]] = {header: [] for header in headers}
        preview: list[list[Any]] = []
        for row in iterator:
            rows += 1
            if len(preview) < 5:
                preview.append([str(value) if value is not None else None for value in row])
            for i, header in enumerate(headers):
                value = row[i] if i < len(row) else None
                if value is None or (isinstance(value, str) and not value.strip()):
                    missing[header] += 1
                elif isinstance(value, (int, float)) and not isinstance(value, bool):
                    numeric[header].append(float(value))
        numeric_summary = {
            header: {"min": min(values), "max": max(values)}
            for header, values in numeric.items()
            if values
        }
        sheets.append(
            {
                "name": worksheet.title,
                "rows": rows,
                "columns": headers,
                "missing_by_column": missing,
                "numeric_ranges": numeric_summary,
                "preview": preview,
                "merged_ranges": [str(value) for value in worksheet.merged_cells.ranges],
            }
        )
    workbook.close()
    return {
        "source_id": source_id,
        "source_uri": source_uri,
        "format": path.suffix.lower().lstrip("."),
        "status": "ok",
        "sheets": sheets,
    }


def _attachment_summary(path: Path, source_id: str, source_uri: str) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        return _csv_summary(path, source_id, source_uri)
    if suffix in {".xlsx", ".xlsm"}:
        return _xlsx_summary(path, source_id, source_uri)
    if suffix in _TEXT_SUFFIXES:
        text = _clean_text(_read_text(path))
        return {
            "source_id": source_id,
            "source_uri": source_uri,
            "format": suffix.lstrip("."),
            "status": "ok",
            "char_count": len(text),
            "preview": text[:1200],
        }
    if suffix == ".pdf":
        pages = _page_records(path, source_id)
        return {
            "source_id": source_id,
            "source_uri": source_uri,
            "format": "pdf",
            "status": "ok" if all(page["status"] == "ok" for page in pages) else "blocked",
            "pages": len(pages),
            "page_char_counts": [page["normalized_char_count"] for page in pages],
        }
    if suffix in {".npy", ".npz"}:
        try:
            import numpy as np

            loaded = np.load(path, allow_pickle=False)
            names = list(loaded.files) if suffix == ".npz" else [path.stem]
            arrays = []
            for name in names:
                array = loaded[name] if suffix == ".npz" else loaded
                arrays.append({"name": name, "shape": list(array.shape), "dtype": str(array.dtype)})
            if hasattr(loaded, "close"):
                loaded.close()
            return {"source_id": source_id, "source_uri": source_uri, "format": suffix[1:], "status": "ok", "arrays": arrays}
        except Exception as exc:  # pragma: no cover - optional binary formats vary
            return {"source_id": source_id, "source_uri": source_uri, "format": suffix[1:], "status": "blocked", "error": str(exc)}
    if suffix in {".png", ".jpg", ".jpeg"}:
        return {
            "source_id": source_id,
            "source_uri": source_uri,
            "format": suffix[1:],
            "status": "review_required",
            "error": "图像附件未自动 OCR/视觉核验，不能作为已审计证据",
        }
    return {
        "source_id": source_id,
        "source_uri": source_uri,
        "format": suffix.lstrip(".") or "unknown",
        "status": "blocked",
        "error": "没有可靠解析器",
    }


def _iter_attachment_files(paths: Iterable[Path], staging: Path) -> Iterator[tuple[Path, str]]:
    for root in paths:
        if root.is_dir():
            for path in sorted(item for item in root.rglob("*") if item.is_file()):
                yield path, str(path)
            continue
        if root.suffix.lower() != ".zip":
            if root.is_file():
                yield root, str(root)
            continue
        with zipfile.ZipFile(root) as archive:
            for index, info in enumerate(archive.infolist(), start=1):
                if info.is_dir() or Path(info.filename).suffix.lower() not in _SUPPORTED_ATTACHMENT_SUFFIXES:
                    continue
                safe_name = Path(info.filename).name or f"entry_{index}"
                target = staging / f"{root.stem}_{index}_{safe_name}"
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(archive.read(info))
                yield target, f"{root.name}::{info.filename}"


def _write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False) + "\n")


def build_evidence_pack(
    problem_path: Path,
    paper_path: Path,
    attachment_paths: Iterable[Path],
    output_dir: Path,
) -> dict[str, Any]:
    """构建证据包；任何无法可靠解析的输入都会让 gate=false。"""
    problem_path = Path(problem_path)
    paper_path = Path(paper_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix="evidence_staging_", dir=str(output_dir)))
    try:
        if not problem_path.exists() or not paper_path.exists():
            missing = [str(path) for path in (problem_path, paper_path) if not path.exists()]
            result = {"schema_version": SCHEMA_VERSION, "status": "blocked", "quality": {"gate": False, "errors": [f"输入不存在：{missing}"]}}
            _write_json(output_dir / "evidence_pack.json", result)
            return result

        sources: list[dict[str, Any]] = []
        source_id = 0

        def register(path: Path, role: str, source_uri: str) -> str:
            nonlocal source_id
            source_id += 1
            current = f"SRC-{source_id:03d}"
            sources.append(
                {
                    "source_id": current,
                    "role": role,
                    "name": path.name,
                    "source_uri": source_uri,
                    "suffix": path.suffix.lower(),
                    "size": path.stat().st_size,
                    "sha256": _sha256(path),
                }
            )
            return current

        problem_source = register(problem_path, "problem", str(problem_path))
        paper_source = register(paper_path, "paper", str(paper_path))
        problem_pages = _page_records(problem_path, problem_source)
        paper_pages = _page_records(paper_path, paper_source)
        paper_chunks = _chunks(paper_pages, paper_source)
        questions = _question_index(problem_path, problem_pages)
        attachment_summaries: list[dict[str, Any]] = []
        attachment_paths = list(attachment_paths)
        for path, source_uri in _iter_attachment_files(attachment_paths, staging):
            current = register(path, "attachment", source_uri)
            attachment_summaries.append(_attachment_summary(path, current, source_uri))

        checks: list[dict[str, Any]] = []

        def check(check_id: str, passed: bool, detail: str) -> None:
            checks.append({"id": check_id, "passed": passed, "detail": detail})

        problem_sparse = [page["page"] for page in problem_pages if page["status"] != "ok"]
        paper_sparse = [page["page"] for page in paper_pages if page["status"] != "ok"]
        check("problem_text", bool(problem_pages) and not problem_sparse, f"问题页数={len(problem_pages)}，稀疏页={problem_sparse}")
        check("problem_questions", bool(questions), f"识别小问={len(questions)}")
        check("paper_text", bool(paper_pages) and not paper_sparse, f"论文页数={len(paper_pages)}，稀疏页={paper_sparse}")
        check("paper_chunks", bool(paper_chunks), f"论文证据片段={len(paper_chunks)}")
        bad_attachments = [item for item in attachment_summaries if item.get("status") != "ok"]
        check("attachment_parsers", not bad_attachments, f"附件总数={len(attachment_summaries)}，阻断/待人工={len(bad_attachments)}")
        gate = all(item["passed"] for item in checks)

        problem_index = {
            "schema_version": SCHEMA_VERSION,
            "source_id": problem_source,
            "source": str(problem_path),
            "pages": len(problem_pages),
            "question_ids": [item["question_id"] for item in questions],
            "questions": questions,
        }
        paper_index = {
            "schema_version": SCHEMA_VERSION,
            "source_id": paper_source,
            "source": str(paper_path),
            "pages": len(paper_pages),
            "chunk_count": len(paper_chunks),
        }
        _write_json(output_dir / "problem_index.json", problem_index)
        _write_json(output_dir / "attachment_manifest.json", {"attachments": attachment_summaries})
        _write_jsonl(output_dir / "problem_pages.jsonl", problem_pages)
        _write_jsonl(output_dir / "paper_pages.jsonl", paper_pages)
        _write_jsonl(output_dir / "paper_chunks.jsonl", paper_chunks)

        question_dir = output_dir / "question_packs"
        question_dir.mkdir(exist_ok=True)
        for question in questions:
            evidence = _retrieve(question, paper_chunks)
            pack = {
                "schema_version": SCHEMA_VERSION,
                "question_id": question["question_id"],
                "problem": {
                    "location": question["location"],
                    "title": question["title"],
                    "text": question["text"],
                },
                "evidence": [
                    {
                        "evidence_id": item["evidence_id"],
                        "location": item["location"],
                        "anchors": item["anchors"],
                        "text": item["text"],
                    }
                    for item in evidence
                ],
                "attachments": attachment_summaries,
                "retrieval": {"count": len(evidence), "method": "deterministic_token_overlap"},
            }
            _write_json(question_dir / f"{question['question_id']}.json", pack)

        result = {
            "schema_version": SCHEMA_VERSION,
            "status": "ready" if gate else "blocked",
            "quality": {
                "gate": gate,
                "checks": checks,
                "errors": [item["detail"] for item in checks if not item["passed"]],
                "warnings": [],
            },
            "sources": sources,
            "problem": problem_index,
            "paper": paper_index,
            "attachments": {"count": len(attachment_summaries)},
            "artifacts": {
                "problem_index": "problem_index.json",
                "paper_pages": "paper_pages.jsonl",
                "paper_chunks": "paper_chunks.jsonl",
                "attachment_manifest": "attachment_manifest.json",
                "question_packs": "question_packs/",
            },
        }
        _write_json(output_dir / "evidence_pack.json", result)
        return result
    finally:
        shutil.rmtree(staging, ignore_errors=True)


def load_evidence_pack(pack_dir: Path) -> dict[str, Any]:
    pack_dir = Path(pack_dir)
    path = pack_dir / "evidence_pack.json"
    if not path.exists():
        raise FileNotFoundError(f"证据包缺少 {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def review_context(pack_dir: Path, max_chars: int = 60_000) -> str:
    """为高级模型提供带定位的证据摘要，不把整篇论文重复塞入提示词。"""
    pack_dir = Path(pack_dir)
    manifest = load_evidence_pack(pack_dir)
    if not manifest["quality"]["gate"]:
        raise ValueError("证据包质量门禁未通过，禁止直接喂给评审模型")
    blocks = [
        f"证据包状态：{manifest['status']}，来源数：{len(manifest['sources'])}",
        "所有引用必须保留 evidence_id 或文件页码；不能把未列出的内容当作事实。",
    ]
    for question_id in manifest["problem"]["question_ids"]:
        path = pack_dir / "question_packs" / f"{question_id}.json"
        item = json.loads(path.read_text(encoding="utf-8"))
        blocks.append(f"\n===== {question_id} · {item['problem']['location']} =====\n{item['problem']['text']}")
        for evidence in item["evidence"]:
            blocks.append(f"[{evidence['evidence_id']} · {evidence['location']}]\n{evidence['text']}")
        blocks.append("附件摘要：" + json.dumps(item["attachments"], ensure_ascii=False))
    return "\n".join(blocks)[:max_chars]

