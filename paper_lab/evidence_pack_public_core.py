"""证据包公开入口与 Excel、小问索引、上下文预算兼容层。"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Iterable

from . import evidence_pack_core as _core


def _xlsx_summary(path, source_id: str, source_uri: str) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {
            "source_id": source_id,
            "source_uri": source_uri,
            "format": path.suffix.lower().lstrip("."),
            "status": "blocked",
            "error": "缺少 openpyxl，不能可靠审计 Excel 附件",
        }
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        iterator = worksheet.iter_rows(values_only=True)
        first = next(iterator, ())
        headers = [
            str(value) if value is not None else f"column_{i + 1}"
            for i, value in enumerate(first)
        ]
        rows = 0
        missing = {header: 0 for header in headers}
        numeric: dict[str, list[float]] = {header: [] for header in headers}
        preview: list[list[Any]] = []
        for row in iterator:
            rows += 1
            if len(preview) < 5:
                preview.append([str(value) if value is not None else None for value in row])
            for i, header in enumerate(headers):
                value = row[i] if i < len(row) else None
                if value is None or (isinstance(value, str) and not value.strip()):
                    missing[header] += 1
                elif isinstance(value, (int, float)) and not isinstance(value, bool):
                    numeric[header].append(float(value))
        numeric_summary = {
            header: {"min": min(values), "max": max(values)}
            for header, values in numeric.items()
            if values
        }
        sheets.append(
            {
                "name": worksheet.title,
                "rows": rows,
                "columns": headers,
                "missing_by_column": missing,
                "numeric_ranges": numeric_summary,
                "preview": preview,
                "merged_ranges": [],
            }
        )
    workbook.close()
    return {
        "source_id": source_id,
        "source_uri": source_uri,
        "format": path.suffix.lower().lstrip("."),
        "status": "ok",
        "sheets": sheets,
    }


def _question_index(problem_path, pages):
    text = "\n".join(page["text"] for page in pages)
    pattern = re.compile(
        r"问题\s*([一二三四五六七八九十0-9]+)|第\s*([一二三四五六七八九十0-9]+)\s*问"
    )
    appendix = re.search(r"(?:^|\n)\s*附录\s*[一二三四五六七八九十0-9]+", text)
    appendix_start = appendix.start() if appendix else len(text)
    accepted = []
    seen: set[int] = set()
    for match in pattern.finditer(text):
        if match.start() >= appendix_start:
            break
        number = _core._question_number(match.group(1) or match.group(2), len(accepted) + 1)
        if number in seen:
            continue
        seen.add(number)
        accepted.append(match)
    questions = []
    for index, match in enumerate(accepted, start=1):
        number = _core._question_number(match.group(1) or match.group(2), index)
        start = match.start()
        next_start = accepted[index].start() if index < len(accepted) else appendix_start
        questions.append(
            {
                "question_id": f"Q{number}",
                "title": match.group(0).strip(),
                "text": text[start:next_start].strip(),
                "location": f"{problem_path.name} chars {start}-{next_start}",
                "source_id": pages[0]["source_id"] if pages else "",
            }
        )
    return questions


def _tokens(text: str) -> set[str]:
    return {
        item.lower()
        for item in re.findall(r"[\u4e00-\u9fff]{2,}|[A-Za-z][A-Za-z0-9_]{1,}|\d+", text)
    }


def _compact_attachment(item: dict[str, Any]) -> dict[str, Any]:
    compact = {
        key: item[key]
        for key in ("source_id", "source_uri", "format", "status", "rows", "char_count", "pages", "arrays")
        if key in item
    }
    if "columns" in item:
        compact["column_count"] = len(item["columns"])
        compact["columns_sample"] = item["columns"][:16]
    if "sheets" in item:
        sheets = []
        for sheet in item["sheets"]:
            missing = {
                key: value
                for key, value in sheet.get("missing_by_column", {}).items()
                if value
            }
            ranges = dict(list(sheet.get("numeric_ranges", {}).items())[:8])
            sheets.append(
                {
                    "name": sheet.get("name"),
                    "rows": sheet.get("rows"),
                    "column_count": len(sheet.get("columns", [])),
                    "columns_sample": sheet.get("columns", [])[:16],
                    "missing_column_count": len(missing),
                    "missing_rows_sample": dict(list(missing.items())[:8]),
                    "numeric_ranges_sample": ranges,
                }
            )
        compact["sheets"] = sheets
    if item.get("preview"):
        compact["preview"] = item["preview"][:2]
    if item.get("error"):
        compact["error"] = item["error"]
    return compact


def _compact_question_packs(output_dir: Path) -> None:
    manifest_path = output_dir / "attachment_manifest.json"
    question_dir = output_dir / "question_packs"
    if not manifest_path.exists() or not question_dir.exists():
        return
    attachments = json.loads(manifest_path.read_text(encoding="utf-8")).get("attachments", [])
    for path in sorted(question_dir.glob("Q*.json")):
        question = json.loads(path.read_text(encoding="utf-8"))
        query = _tokens(question["problem"]["text"])
        ranked = []
        for item in attachments:
            score = len(query & _tokens(json.dumps(item, ensure_ascii=False)))
            ranked.append((score, item.get("source_id", ""), item))
        ranked.sort(key=lambda value: (-value[0], value[1]))
        selected = [item for score, _, item in ranked if score > 0][:8] or [item for _, _, item in ranked[:8]]
        question["attachments"] = [_compact_attachment(item) for item in selected]
        question["attachments_index"] = {
            "manifest": "attachment_manifest.json",
            "selection": "top_8_deterministic_token_overlap",
            "total_available": len(attachments),
        }
        path.write_text(json.dumps(question, ensure_ascii=False, indent=2), encoding="utf-8")


def build_evidence_pack(problem_path: Path, paper_path: Path, attachment_paths: Iterable[Path], output_dir: Path):
    result = _core.build_evidence_pack(problem_path, paper_path, attachment_paths, output_dir)
    _compact_question_packs(Path(output_dir))
    return result


_core._xlsx_summary = _xlsx_summary
_core._question_index = _question_index

load_evidence_pack = _core.load_evidence_pack
review_context = _core.review_context

__all__ = ["build_evidence_pack", "load_evidence_pack", "review_context"]

